import os
import re
import pandas as pd
from PIL import Image
import pytesseract
from docx import Document

# Point to Tesseract executable on Windows if you get a "tesseract is not installed" error:
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

def clean_text(text: str) -> str:
    """Normalize whitespace and strip leading/trailing spaces."""
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n+', '\n', text)
    return text.strip()

def extract_from_docx(file_path: str) -> str:
    """Extract text from paragraphs and tables in a .docx file."""
    doc = Document(file_path)
    full_text = []
    
    # Extract paragraph text
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
            
    # Extract text from tables if any exist
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
                
    return clean_text("\n".join(full_text))

def extract_from_image(file_path: str) -> str:
    """Extract text from scanned PNG using OCR."""
    img = Image.open(file_path).convert("L")  # Convert to grayscale for better OCR
    raw_text = pytesseract.image_to_string(img)
    return clean_text(raw_text)

def parse_document(file_path: str) -> str:
    """Route document to appropriate parser based on extension."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return extract_from_docx(file_path)
    elif ext in [".png", ".jpg", ".jpeg"]:
        return extract_from_image(file_path)
    else:
        raise ValueError(f"Unsupported format: {ext}")

def load_and_prepare_dataset(csv_path: str, data_dir: str):
    df = pd.read_csv(csv_path)
    extracted_texts = []
    
    # Locate the filename column in your train.csv (adjust column name if needed)
    filename_col = df.columns[0]
    
    for _, row in df.iterrows():
        base_name = str(row[filename_col]).strip()
        
        # Match against files in data_dir (handling extensions)
        matched_file = None
        for fname in os.listdir(data_dir):
            if fname.startswith(base_name) or base_name in fname:
                matched_file = os.path.join(data_dir, fname)
                break
                
        if matched_file and os.path.exists(matched_file):
            print(f"Extracting: {matched_file}")
            text = parse_document(matched_file)
            extracted_texts.append(text)
        else:
            print(f"Warning: File not found for {base_name}")
            extracted_texts.append("")
            
    df['document_text'] = extracted_texts
    return df

if __name__ == "__main__":
    train_csv = os.path.join("data", "train.csv")
    train_dir = os.path.join("data", "train")
    
    df_train = load_and_prepare_dataset(train_csv, train_dir)
    df_train.to_json("processed_train.json", orient="records", indent=2)
    print("Extraction complete. Processed data saved to processed_train.json.")